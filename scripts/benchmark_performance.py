"""Reproducible before/after benchmarks; isolated in-memory SQLite, no network.

Run with .venv/bin/python scripts/benchmark_performance.py desktop|cloud|matrix [baseline].
Times cover the named operation, not browser/network latency. SQL counts are
measured separately so query logging does not distort the timing samples.
"""
import copy
import inspect
import io
import json
import os
import platform
import random
import statistics
import subprocess
import sys
import time
import types
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
BASE = sys.argv[2] if len(sys.argv) > 2 else '366a37a'
MODE = sys.argv[1]
sys.path.insert(0, str(ROOT / 'backend' if MODE == 'cloud' else ROOT))
os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
from django.conf import settings
settings.DATABASES['default']['NAME'] = ':memory:'
settings.DEBUG = False
import django
django.setup()
from django.core.management import call_command
from django.db import connection
from django.test import RequestFactory
from django.test.utils import CaptureQueriesContext
call_command('migrate', verbosity=0)


def old_module(path, name):
    source = subprocess.check_output(['git', 'show', f'{BASE}:{path}'], cwd=ROOT, text=True)
    module = types.ModuleType(name)
    module.__file__ = str(ROOT / path)
    module.__package__ = name.rsplit('.', 1)[0]
    exec(compile(source, module.__file__, 'exec'), module.__dict__)
    return module


def bench(label, before, after, validate=None, rounds=9):
    a, b = before(), after()
    if validate:
        validate(a, b)
    samples = [[], []]
    for index in range(rounds):
        for which in ((0, 1) if index % 2 == 0 else (1, 0)):
            start = time.perf_counter_ns()
            (before, after)[which]()
            samples[which].append((time.perf_counter_ns() - start) / 1e6)
    counts = []
    for fn in (before, after):
        with CaptureQueriesContext(connection) as queries:
            fn()
        counts.append(len(queries))
    medians = [statistics.median(sample) for sample in samples]
    print(json.dumps(dict(case=label, before_ms=round(medians[0], 4),
        after_ms=round(medians[1], 4), speedup=round(medians[0]/medians[1], 3),
        reduction_pct=round(100*(1-medians[1]/medians[0]), 2),
        queries=counts, samples_ms=samples), ensure_ascii=False), flush=True)


def equal(a, b):
    assert a == b


print(json.dumps(dict(mode=MODE, baseline=BASE, optimized=subprocess.check_output(
    ['git', 'rev-parse', '--short', 'HEAD'], cwd=ROOT, text=True).strip(),
    python=platform.python_version(), platform=platform.platform(),
    optimized_note='working tree; same inputs and configuration on both sides',
    database='SQLite :memory:', rounds=9, warmups=1)), flush=True)

if MODE == 'matrix':
    from seats import sorting, views
    from seats.models import Classroom, Student, LayoutSnapshot
    old_sort = old_module('seats/sorting.py', 'seats._benchmark_matrix_sort')
    old_views = old_module('seats/views.py', 'seats._benchmark_matrix_views')
    for n in (60, 120, 300, 600, 1200):
        rng = random.Random(20260924)
        students = [types.SimpleNamespace(pk=i+1, name=rng.choice(['张','李','陈','刘','王'])+str(i),
            score=rng.randrange(101),student_id=str(i),custom_data={}) for i in range(n)]
        rng.shuffle(students)
        for mode, rule in (
            ('pinyin', {'rules':[{'field':'name','transform':'pinyin','direction':'asc'}]}),
            ('score', {'rules':[{'field':'score','transform':'numeric','direction':'desc'}]}),
            ('mixed', {'rules':[{'field':'score','transform':'numeric','direction':'desc'},
                {'field':'name','transform':'pinyin','direction':'asc'}]}),
        ):
            bench(f'matrix_sort_{mode}_{n}',lambda:old_sort.sort_students(students,rule),
                lambda:sorting.sort_students(students,rule),equal)
    classroom=Classroom.objects.create(name='快照增长测试',rows=6,cols=10)
    Student.objects.bulk_create([Student(classroom=classroom,name=f'学生{i}',score=i%100) for i in range(60)])
    payload=views._serialize_seats_file_bundle(classroom)
    for n in (1, 10, 30, 100):
        classroom.layout_snapshots.all().delete()
        LayoutSnapshot.objects.bulk_create([LayoutSnapshot(classroom=classroom,name=f'快照{i}',data=payload) for i in range(n)])
        def snapshot_rows(optimized):
            qs=classroom.layout_snapshots.all()
            if optimized: qs=qs.only('pk','classroom_id','name','created_at')
            return [(row.pk,row.name,row.created_at) for row in qs]
        bench(f'matrix_snapshot_count_{n}',lambda:snapshot_rows(False),lambda:snapshot_rows(True),equal)

elif MODE == 'desktop':
    import openpyxl
    from seats import sorting, constraints, views
    from seats.models import (Classroom, Student, Seat, SeatGroup, SeatConstraint,
        LayoutSnapshot, ClassroomHistoryEntry, AIConversation, AIConversationMessage)
    old_sort = old_module('seats/sorting.py', 'seats._benchmark_sort')
    old_constraints = old_module('seats/constraints.py', 'seats._benchmark_constraints')
    old_views = old_module('seats/views.py', 'seats._benchmark_views')
    for n, history_count, snapshot_count in ((60, 20, 10), (600, 300, 100)):
        rng = random.Random(20260924)
        classroom = Classroom.objects.create(name=f'基准{n}', rows=n//10, cols=10)
        classroom.seats.all().delete()  # Classroom creation initializes an empty grid.
        students = Student.objects.bulk_create([
            Student(classroom=classroom, name=rng.choice(['张','李','王','陈','刘']) + str(i),
                score=rng.randrange(101), student_id=str(i)) for i in range(n)])
        groups = SeatGroup.objects.bulk_create([
            SeatGroup(classroom=classroom, name=f'第{i}组', order=i) for i in range(n//6)])
        Seat.objects.bulk_create([Seat(classroom=classroom, student=s,
            row=i//10+1, col=i%10+1, group=groups[i//6]) for i,s in enumerate(students)])
        rng.shuffle(students)
        rule = {'rules':[{'field':'name','transform':'pinyin','direction':'asc'}]}
        bench(f'sort_pinyin_{n}', lambda:old_sort.sort_students(students,rule),
            lambda:sorting.sort_students(students,rule), equal)
        constraint_count = n//3
        SeatConstraint.objects.bulk_create([SeatConstraint(classroom=classroom,
            student=s, constraint_type='must_row', row=i%(n//10)+1)
            for i,s in enumerate(students[:constraint_count])])
        bench(f'constraint_diagnostics_{n}_{constraint_count}',
            lambda:old_constraints.serialize_constraints(classroom),
            lambda:constraints.serialize_constraints(classroom), equal)
        request = RequestFactory().get('/')
        def report_valid(a,b):
            def cells(response):
                wb = openpyxl.load_workbook(io.BytesIO(response.content))
                return list(wb.active.values)
            equal(cells(a),cells(b))
        bench(f'group_report_{n}_{len(groups)}',
            lambda:old_views.export_group_report(request,classroom.pk),
            lambda:views.export_group_report(request,classroom.pk), report_valid)
        base = views._serialize_seats_file_bundle(classroom)
        ClassroomHistoryEntry.objects.bulk_create([ClassroomHistoryEntry(
            classroom=classroom, action_type='move', payload={'before':base['current_state'],
            'after':base['current_state']}) for _ in range(history_count)])
        conversation = AIConversation.objects.create(classroom=classroom)
        AIConversationMessage.objects.bulk_create([AIConversationMessage(
            conversation=conversation, content='模拟排座建议。'*100) for _ in range(history_count)])
        session = types.SimpleNamespace(limits={'max_history_steps':0,'sync_ai_conversations':False})
        def export_valid(a,b):
            a=copy.deepcopy(a); b=copy.deepcopy(b)
            for item in (a,b): item.get('meta',{}).pop('exported_at',None)
            equal(a,b)
        bench(f'cloud_export_history_disabled_{n}_{history_count}',
            lambda:old_views._cloud_export_payload(classroom,session),
            lambda:views._cloud_export_payload(classroom,session), export_valid)
        LayoutSnapshot.objects.bulk_create([LayoutSnapshot(classroom=classroom,
            name=f'快照{i}',data=base) for i in range(snapshot_count)])
        def snapshots(optimized):
            qs=classroom.layout_snapshots.all()
            if optimized: qs=qs.only('pk','classroom_id','name','created_at')
            return [(x.pk,x.name,x.created_at) for x in qs]
        bench(f'snapshot_list_query_{snapshot_count}_{len(json.dumps(base))}_chars',
            lambda:snapshots(False),lambda:snapshots(True),equal)
        wb=openpyxl.Workbook()
        for row in range(n//10): wb.active.append([f'学生{row*10+col}' for col in range(10)])
        stream=io.BytesIO(); wb.save(stream); blob=stream.getvalue(); wb.close()
        options=dict(manual_name_terms=[],manual_podium_terms=[],manual_empty_terms=[],
            manual_aisle_terms=[],auto_detect_names=True,layout_transform='none')
        def excel(module, optimized):
            stream=io.BytesIO(blob)
            if not optimized:
                defaults=module._detect_layout_import_defaults(stream,options)
                opts={**options,'layout_transform':defaults['layout_transform']}
                return module._build_layout_preview_response(stream,defaults['start_row'],defaults['end_row'],opts)
            book=openpyxl.load_workbook(stream,data_only=True)
            try:
                bounds=module._detect_layout_bounds(book.active)
                defaults=module._detect_layout_import_defaults(stream,options,workbook=book,bounds=bounds)
                opts={**options,'layout_transform':defaults['layout_transform']}
                return module._build_layout_preview_response(stream,defaults['start_row'],defaults['end_row'],opts,workbook=book,bounds=bounds)
            finally: book.close()
        bench(f'excel_parse_preview_{n}',lambda:excel(old_views,False),lambda:excel(views,True),equal)

elif MODE == 'cloud':
    from django.http import JsonResponse
    from cloud import views, sync
    from cloud.models import CloudUser, CloudClassroom, CloudClassroomCommit, CloudSnapshot
    from cloud.delta import build_delta, apply_delta, snapshot_hash
    from cloud.tests import CloudSyncStorageTests
    old_views=old_module('backend/cloud/views.py','cloud._benchmark_views')
    old_sync=old_module('backend/cloud/sync.py','cloud._benchmark_sync')
    user=CloudUser.objects.create(uid='isolated-benchmark')
    # Measure endpoint database + JSON work; omit auth/encryption equally on both sides.
    def response(payload, **kwargs): return JsonResponse(payload)
    views._encrypted_json_response=old_views._encrypted_json_response=response
    for n, history_count, snapshot_count in ((60,20,10),(600,10,50)):
        data=CloudSyncStorageTests().fuckseats_snapshot(rows=n//10,cols=10)
        template=data['students'][0]
        data['students']=[{**copy.deepcopy(template),'name':f'学生{i}','student_id':str(i)} for i in range(n)]
        data['current_state']['students']=[{**copy.deepcopy(x),'pk':i+1} for i,x in enumerate(data['students'])]
        layout_snapshot=copy.deepcopy(data)
        data['history']['entries']=[{'action_type':'move','payload':{'before':copy.deepcopy(data['current_state']),
            'after':copy.deepcopy(data['current_state'])}} for _ in range(history_count)]
        data['ai_conversations']=[{'title':'模拟对话','messages':[{'role':'user','content':'建议'*100,'payload':{}}]}]
        sync.validate_fuckseats_snapshot(data)
        sync.check_payload_size(data)  # Respect the default server upload limit.
        classroom=CloudClassroom.objects.create(user=user,name='基准班',rows=n//10,cols=10,data_snapshot=data,version=1)
        CloudSnapshot.objects.bulk_create([CloudSnapshot(classroom=classroom,
            name=f'快照{i}',data=layout_snapshot,size_bytes=len(json.dumps(layout_snapshot))) for i in range(snapshot_count)])
        req=RequestFactory().get('/');req.cloud_user=user;req.cloud_session=None
        for name,args in [('sync_status',()),('snapshots_list',(classroom.uuid,))]:
            bench(f'{name}_{n}_{history_count}',
                lambda name=name,args=args:inspect.unwrap(getattr(old_views,name))(req,*args),
                lambda name=name,args=args:inspect.unwrap(getattr(views,name))(req,*args))
        changed=copy.deepcopy(data)
        changed['students'][0]['score']=91
        changed['current_state']['students'][0]['score']=91
        changed['meta']['exported_at']='2026-09-24T12:00:00+00:00'
        changes=build_delta(data,changed)
        assert apply_delta(data,changes)==changed
        classroom.data_snapshot=changed;classroom.version=2;classroom.save()
        CloudClassroomCommit.objects.create(classroom=classroom,version=2,base_version=1,
            changes=changes,snapshot_hash=snapshot_hash(changed))
        delta_req=RequestFactory().get('/?protocol=delta-v1&base_version=1')
        delta_req.cloud_user=user;delta_req.cloud_session=None
        full=lambda:inspect.unwrap(old_views.sync_pull)(req,classroom.uuid)
        delta=lambda:inspect.unwrap(views.sync_pull)(delta_req,classroom.uuid)
        bench(f'sync_pull_one_change_{n}_{history_count}',full,delta)
        envelope={'uuid':str(classroom.uuid),'base_version':1,'cloud_version':1,'local_version':2,
            'force':False,'device_id':'benchmark','last_operation_at':'2026-09-24T12:00:00+00:00'}
        full_push={**envelope,'data':changed}
        delta_push={**envelope,'protocol':'delta-v1','changes':changes,'snapshot_hash':snapshot_hash(changed)}
        size=lambda p:len(json.dumps(p,ensure_ascii=False,separators=(',',':')).encode())
        print(json.dumps(dict(case=f'wire_json_bytes_{n}_{history_count}',
            full_push=size(full_push),delta_push=size(delta_push),
            full_pull=len(full().content),delta_pull=len(delta().content)),ensure_ascii=False),flush=True)
        with_history=copy.deepcopy(changed)
        with_history['history']['entries'].append({'action_type':'score','payload':{
            'before':data['current_state'],'after':changed['current_state']}})
        history_delta=build_delta(data,with_history)
        assert apply_delta(data,history_delta)==with_history
        print(json.dumps(dict(case=f'wire_json_score_and_history_append_{n}_{history_count}',
            full_push=size({**envelope,'data':with_history}),
            delta_push=size({**envelope,'protocol':'delta-v1','changes':history_delta,
                'snapshot_hash':snapshot_hash(with_history)})),ensure_ascii=False),flush=True)
        bench(f'child_records_unchanged_{n}_{history_count}',
            lambda:old_sync.sync_child_records(classroom,changed),
            lambda:sync.sync_child_records(classroom,changed,sections={'students','current_state','meta'}))
        # Keep each scenario independent for status queries.
        classroom.delete()
